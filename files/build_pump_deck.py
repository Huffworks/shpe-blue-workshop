"""Build the pre-generated Kestrel pump briefing decks for the Blue workshop demo.

This script is the fallback for demo beats 2 and 3. It reads the Tarrant brand
values from the "## Values" section of vault/brand-kit.md, builds a 5-slide
deck from the Kestrel pump vibration test memo that follows the slide
rules, reloads the result with python-pptx and checks the rules it can check.

Run it from anywhere:

    .venv/bin/python demo/outputs/build_pump_deck.py              # both decks
    .venv/bin/python demo/outputs/build_pump_deck.py --ask-first  # only the variant

The default deck puts the ask on the last slide. The --ask-first variant is the
deck after the memory line "Exec audiences want the ask on slide 1, not slide 5."
"""

import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
BRAND_KIT = ROOT / "vault" / "brand-kit.md"
MEMO = ROOT / "vault" / "docs" / "kestrel-pump-vibration-test-memo.md"
MEMO_NAME = "Kestrel pump skid vibration test memo, 2026-10-08"
FOOTER_TEXT = "Tarrant Dynamics | Internal"

MAX_BULLETS = 5
MAX_WORDS = 12


def read_brand(path):
    """Return the "- Key: value" pairs from the "## Values" section."""
    values = {}
    in_values = False
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            in_values = line.strip() == "## Values"
            continue
        match = re.match(r"^- ([^:]+): (.+)$", line)
        if in_values and match:
            values[match.group(1).strip()] = match.group(2).strip()
    required = ["Company", "Primary", "Secondary", "Accent", "Background",
                "Light", "Heading font", "Body font"]
    missing = [key for key in required if key not in values]
    if missing:
        sys.exit(f"brand-kit.md is missing these values: {', '.join(missing)}")
    return values


def rgb(hex_value):
    return RGBColor.from_string(hex_value.lstrip("#").upper())


# A bullet is a list of (text, highlighted) runs. At most one run per slide is
# highlighted, and it is drawn in the accent colour (slide rule 5).
EVIDENCE_SLIDES = [
    {
        "title": "Pump vibration exceeds the limit only at the top test speed",
        "source": "Results table",
        "bullets": [
            [("Pump drive end: ", False), ("6.8 mm/s RMS", True), (" at 3,550 rpm.", False)],
            [("Program limit is 4.5 mm/s RMS at any bearing housing.", False)],
            [("Motor drive end is also over, at 4.6 mm/s RMS.", False)],
            [("All bearings are under the limit at 1,800 to 3,000 rpm.", False)],
            [("Both non-drive-end readings stay under the limit at 3,550 rpm.", False)],
        ],
    },
    {
        "title": "The spectrum points to alignment or the coupling, not hydraulics",
        "source": "Spectrum at 3,550 rpm table",
        "bullets": [
            [("2x running speed peak: ", False), ("5.9 mm/s RMS", True), (" at 118.3 Hz.", False)],
            [("1x peak is 2.1 mm/s RMS, which argues against imbalance.", False)],
            [("Vane pass peak is 0.8 mm/s RMS, which argues against hydraulics.", False)],
            [("A high 2x peak suggests misalignment or a coupling not absorbing it.", False)],
            [("Operating point was 176 m3/h at 41.2 m head.", False)],
        ],
    },
    {
        "title": "Cold alignment was in tolerance, so the root cause is still open",
        "source": "Method and Spectrum sections",
        "bullets": [
            [("Parallel offset was 0.04 mm against a 0.10 mm tolerance.", False)],
            [("Angular offset was 0.03 mm/100 mm against a 0.05 mm/100 mm tolerance.", False)],
            [("Pump drive end climbed 3.9 to 6.8 mm/s RMS above 3,000 rpm.", False)],
            [("Suspects: motor frame thermal growth or VC-38 disc pack wear.", False)],
            [("Flag: ", False), ("no hot alignment reading was taken.", True)],
        ],
    },
]

NEXT_STEPS = [
    [("Install the replacement VC-38 coupling, expected 2026-10-12.", False)],
    [("Repeat the four-speed test from 1,800 to 3,550 rpm.", False)],
    [("Take a hot alignment reading right after the 3,550 rpm hold.", False)],
    [("Flag: the memo predicts the new coupling fixes it, untested.", False)],
]

ASK = "Hold the 72-hour endurance run until the coupling retest passes"


def slide_plan(ask_first):
    """Return the five slides as dicts. Slide 1 is always the title slide."""
    if ask_first:
        title_slide = {
            "kind": "title",
            "title": f"Decision needed: {ASK[0].lower()}{ASK[1:]}",
            "subtitle": (
                "The pump drive end reads 6.8 mm/s RMS at 3,550 rpm against a "
                "4.5 mm/s RMS limit, and the fix is not yet proven."
            ),
            "highlight": "6.8 mm/s RMS",
        }
        closing = {
            "kind": "content",
            "title": "Retest with the new coupling before the 72-hour run starts",
            "source": "Next steps",
            "bullets": NEXT_STEPS + [
                [("Start the 72-hour run only if both drive ends pass.", False)],
            ],
        }
    else:
        title_slide = {
            "kind": "title",
            "title": "The Kestrel pump skid fails the vibration limit only at 3,550 rpm",
            "subtitle": "Briefing on test KT-114 for the Kestrel program team.",
            "highlight": None,
        }
        closing = {
            "kind": "content",
            "title": "Retest with the new coupling before the 72-hour run starts",
            "source": "Next steps",
            "bullets": NEXT_STEPS + [
                [("Decision needed: ", False), ("hold the endurance run until retest passes.", True)],
            ],
        }
    content = [dict(slide, kind="content") for slide in EVIDENCE_SLIDES]
    return [title_slide] + content + [closing]


def add_text(slide, left, top, width, height, text, font, size, color,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def build(brand, ask_first, out_path):
    primary, secondary = rgb(brand["Primary"]), rgb(brand["Secondary"])
    accent, light = rgb(brand["Accent"]), rgb(brand["Light"])
    white = rgb(brand["Background"])
    heading, body = brand["Heading font"], brand["Body font"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    width = prs.slide_width

    for number, spec in enumerate(slide_plan(ask_first), start=1):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = white

        if spec["kind"] == "title":
            band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.2), width, Inches(2.6))
            band.fill.solid()
            band.fill.fore_color.rgb = primary
            band.line.fill.background()
            add_text(slide, Inches(0.8), Inches(2.4), width - Inches(1.6), Inches(2.2),
                     spec["title"], heading, 36, white, bold=True, anchor=MSO_ANCHOR.MIDDLE)
            sub = add_text(slide, Inches(0.8), Inches(5.1), width - Inches(1.6), Inches(1.0),
                           "", body, 20, secondary)
            para = sub.text_frame.paragraphs[0]
            para.runs[0].text = ""
            text, mark = spec["subtitle"], spec["highlight"]
            parts = text.split(mark, 1) if mark else [text]
            for i, part in enumerate(parts):
                if i == 1:
                    run = para.add_run()
                    run.text = mark
                    run.font.name, run.font.size, run.font.bold = body, Pt(20), True
                    run.font.color.rgb = accent
                run = para.add_run()
                run.text = part
                run.font.name, run.font.size = body, Pt(20)
                run.font.color.rgb = secondary
            add_text(slide, Inches(0.8), Inches(6.2), width - Inches(1.6), Inches(0.6),
                     f"Sam Ortega, Systems Engineering. Source: {MEMO_NAME}.",
                     body, 14, secondary)
            continue

        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, width, Inches(1.3))
        bar.fill.solid()
        bar.fill.fore_color.rgb = primary
        bar.line.fill.background()
        add_text(slide, Inches(0.6), Inches(0.15), width - Inches(1.2), Inches(1.0),
                 spec["title"], heading, 28, white, bold=True, anchor=MSO_ANCHOR.MIDDLE)

        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), width - Inches(1.6), Inches(4.4))
        frame = box.text_frame
        frame.word_wrap = True
        for i, runs in enumerate(spec["bullets"]):
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.space_after = Pt(14)
            lead = para.add_run()
            lead.text = "•  "
            lead.font.name, lead.font.size = body, Pt(22)
            lead.font.color.rgb = primary
            for text, highlighted in runs:
                run = para.add_run()
                run.text = text
                run.font.name, run.font.size = body, Pt(22)
                run.font.bold = highlighted
                run.font.color.rgb = accent if highlighted else secondary

        add_text(slide, Inches(0.8), Inches(6.2), width - Inches(1.6), Inches(0.4),
                 f"Source: {MEMO_NAME}, {spec['source']}.", body, 12, secondary)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(6.75),
                                      width - Inches(1.2), Pt(1))
        rule.fill.solid()
        rule.fill.fore_color.rgb = light
        rule.line.fill.background()
        add_text(slide, Inches(0.6), Inches(6.85), Inches(8), Inches(0.4),
                 FOOTER_TEXT, heading, 12, secondary)
        add_text(slide, width - Inches(1.6), Inches(6.85), Inches(1.0), Inches(0.4),
                 str(number), heading, 12, secondary, align=PP_ALIGN.RIGHT)

    prs.save(out_path)


NUMBER = re.compile(r"\d[\d,]*(?:\.\d+)?")


def numbers_in(text):
    return {n.replace(",", "") for n in NUMBER.findall(text)}


def check(path, brand, ask_first):
    """Reload the deck and check the slide rules that can be checked by code."""
    prs = Presentation(path)
    memo_numbers = numbers_in(MEMO.read_text(encoding="utf-8"))
    accent = brand["Accent"].lstrip("#").upper()
    problems = []

    if len(prs.slides) != 5:
        problems.append(f"expected 5 slides, found {len(prs.slides)}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        problems.append("slides are not 16:9")

    for number, slide in enumerate(prs.slides, start=1):
        texts = [s.text_frame for s in slide.shapes if s.has_text_frame and s.text_frame.text]
        all_text = " ".join(f.text for f in texts)
        amber_runs = [r for f in texts for p in f.paragraphs for r in p.runs
                      if r.font.color and r.font.color.type is not None
                      and str(r.font.color.rgb) == accent]
        if len(amber_runs) > 1:
            problems.append(f"slide {number}: accent colour used {len(amber_runs)} times")
        stray = numbers_in(all_text) - memo_numbers - {str(number)}
        if stray:
            problems.append(f"slide {number}: numbers not in the memo: {sorted(stray)}")
        if "Source:" not in all_text:
            problems.append(f"slide {number}: no source named")

        bullets = [p for f in texts for p in f.paragraphs if p.text.startswith("•")]
        if number == 1:
            if bullets:
                problems.append("slide 1: the title slide should have no bullets")
            if FOOTER_TEXT in all_text:
                problems.append("slide 1: the title slide should have no footer")
            if ask_first and "Decision needed" not in all_text:
                problems.append("slide 1: the ask-first deck does not put the ask on slide 1")
            continue
        if len(bullets) > MAX_BULLETS:
            problems.append(f"slide {number}: {len(bullets)} bullets")
        for para in bullets:
            words = len(para.text.lstrip("• ").split())
            if words > MAX_WORDS:
                problems.append(f"slide {number}: {words}-word bullet: {para.text}")
        if FOOTER_TEXT not in all_text or str(number) not in [f.text for f in texts]:
            problems.append(f"slide {number}: footer or slide number missing")
        if not ask_first and number == 5 and "Decision needed" not in all_text:
            problems.append("slide 5: the default deck should carry the ask on the last slide")
        if ask_first and "Decision needed" in all_text:
            problems.append(f"slide {number}: the ask-first deck repeats the ask after slide 1")
    return problems


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--ask-first", action="store_true",
                        help="build only the variant with the ask on slide 1")
    args = parser.parse_args()

    brand = read_brand(BRAND_KIT)
    variants = [True] if args.ask_first else [False, True]
    failed = False
    for ask_first in variants:
        name = "kestrel-pump-briefing-ask-first.pptx" if ask_first else "kestrel-pump-briefing.pptx"
        out_path = HERE / name
        build(brand, ask_first, out_path)
        problems = check(out_path, brand, ask_first)
        if problems:
            failed = True
            print(f"{name}: FAILED")
            for problem in problems:
                print(f"  {problem}")
        else:
            print(f"{name}: built and checked, 5 slides, all rule checks passed")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
